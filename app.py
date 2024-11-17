import os
import logging
from flask import Flask, request, jsonify, send_file
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import random
from groq import Groq
from functools import wraps
from dotenv import load_dotenv
import time
from datetime import datetime
import io
from report import report_bp, generate_feedback_summary
from real_time_feedback import real_time_feedback_bp
from feedback_processor import feedback_processor_bp

# Load environment variables from .env file
load_dotenv()

app = Flask(__name__)

# Register blueprints
app.register_blueprint(report_bp, url_prefix='/report')
app.register_blueprint(real_time_feedback_bp, url_prefix="/real_time_feedback")
app.register_blueprint(feedback_processor_bp, url_prefix='/feedback_processor')

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

# Configure Groq API
GROQ_API_KEY = os.getenv("GROQ_API_KEY")
client = Groq(api_key=GROQ_API_KEY)

def create_feedback_document(data, summary):
    """
    Create a professionally formatted Word document with the feedback.
    """
    doc = Document()
    
    # Document title
    title = doc.add_heading('Pitch Feedback Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Team information
    doc.add_paragraph()
    team_info = doc.add_paragraph()
    team_info.add_run('Team: ').bold = True
    team_info.add_run(data.get('teamName', 'Not specified'))
    team_info.add_run('\nPitch Number: ').bold = True
    team_info.add_run(str(data.get('pitchNumber', 'Not specified')))
    team_info.add_run('\nSession: ').bold = True
    team_info.add_run(str(data.get('session', 'Not specified')))
    team_info.add_run('\nDate: ').bold = True
    team_info.add_run(datetime.now().strftime('%B %d, %Y'))
    
    # Executive Summary
    doc.add_heading('Executive Summary', 1)
    doc.add_paragraph(summary)
    
    # Detailed Scoring
    doc.add_heading('Detailed Evaluation', 1)
    
    for section in data.get('scoringSections', []):
        # Section heading
        section_heading = doc.add_heading(section['title'], 2)
        
        # Score and weight information
        score_info = doc.add_paragraph()
        score_info.add_run('Score: ').bold = True
        score_info.add_run(f"{section['score']:.1f}/10")
        score_info.add_run('\nWeight: ').bold = True
        score_info.add_run(f"{section['weight']}%")
        
        # Question scores
        if section.get('questionScores'):
            scores_table = doc.add_table(rows=1, cols=2)
            scores_table.style = 'Table Grid'
            header_cells = scores_table.rows[0].cells
            header_cells[0].text = 'Question'
            header_cells[1].text = 'Score'
            
            for question, score in section['questionScores'].items():
                row_cells = scores_table.add_row().cells
                row_cells[0].text = f"Question {int(question) + 1}"
                row_cells[1].text = f"{score}/10"
        
        # Section feedback
        if section.get('feedback'):
            doc.add_paragraph()
            feedback_para = doc.add_paragraph()
            feedback_para.add_run('Feedback: ').bold = True
            feedback_para.add_run(section['feedback'])
        
        doc.add_paragraph()  # Add spacing between sections
    
    # General Feedback
    if data.get('generalFeedback'):
        doc.add_heading('General Feedback', 1)
        doc.add_paragraph(data['generalFeedback'])
    
    # Format the document
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(11)
            run.font.name = 'Calibri'
    
    # Save to memory stream
    doc_stream = io.BytesIO()
    doc.save(doc_stream)
    doc_stream.seek(0)
    
    return doc_stream

def retry_on_quota_exceeded():
    """Decorator to handle quota exceeded errors with exponential backoff"""
    def decorator(func):
        @wraps(func)
        def wrapper(*args, **kwargs):
            max_retries = 5
            base_delay = 1
            
            for attempt in range(max_retries):
                try:
                    return func(*args, **kwargs)
                except Exception as e:
                    if "quota exceeded" in str(e).lower() and attempt < max_retries - 1:
                        delay = (base_delay * 2 ** attempt) + (random.uniform(0, 1))
                        logger.warning(f"Quota exceeded. Retrying in {delay:.2f} seconds. Attempt {attempt + 1}/{max_retries}")
                        time.sleep(delay)
                    else:
                        raise
            
            raise Exception("Max retries exceeded")
        return wrapper
    return decorator

@retry_on_quota_exceeded()
def generate_summary(feedback_data):
    """
    Generate a comprehensive summary using Groq's chat API.
    """
    # Prepare structured prompt
    sections_text = []
    for section in feedback_data.get('scoringSections', []):
        sections_text.append(f"\n{section['title']} (Score: {section['score']:.1f}/10):\n{section.get('feedback', 'No feedback provided')}")
    
    prompt = f"""
    Please provide a professional and constructive executive summary of the following pitch feedback:
    
    Team: {feedback_data.get('teamName')}
    Pitch Number: {feedback_data.get('pitchNumber')}
    Session: {feedback_data.get('session')}
    
    Detailed Feedback:
    {"".join(sections_text)}
    
    General Feedback:
    {feedback_data.get('generalFeedback', 'No general feedback provided')}
    
    Please structure the summary to include:
    1. Overall assessment
    2. Key strengths
    3. Areas for improvement
    4. Strategic recommendations
    
    Keep the tone constructive and professional.
    """
    
    try:
        chat_completion = client.chat.completions.create(
            messages=[
                {"role": "user", "content": prompt}
            ],
            model="llama3-8b-8192"  
        )
        
        summary = chat_completion.choices[0].message.content
        
        if not summary:
            raise ValueError("Empty response received from Groq")
        
        return summary
    
    except Exception as e:
        logger.error(f"Error during summarization: {str(e)}", exc_info=True)
        raise

@app.route("/summarize_feedback", methods=["POST"])
def summarize_feedback():
    """
    API endpoint to process feedback and generate a downloadable Word document.
    """
    try:
        # Get JSON data
        data = request.get_json()
        if not data:
            return jsonify({"error": "No data provided"}), 400
        
        # Generate summary
        try:
            summary = generate_summary(data)
            
            # Create document
            doc_stream = create_feedback_document(data, summary)
            
            # Generate filename
            filename = f"{data.get('teamName', 'Team')}_{data.get('pitchNumber', 'Pitch')}_Feedback.docx"
            filename = "".join(c for c in filename if c.isalnum() or c in ('-', '_')).rstrip()
            
            # Return document
            return send_file(
                doc_stream,
                mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                as_attachment=True,
                download_name=filename
            )
            
        except Exception as e:
            logger.error(f"Error processing feedback: {str(e)}", exc_info=True)
            return jsonify({
                "error": "An error occurred while processing the feedback",
                "details": str(e)
            }), 500
            
    except Exception as e:
        logger.error(f"Unexpected error: {str(e)}", exc_info=True)
        return jsonify({
            "error": "An unexpected error occurred",
            "details": str(e) if app.debug else None
        }), 500

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)